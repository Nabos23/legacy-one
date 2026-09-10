"""Process a user-uploaded chat attachment into something the agent can consume.

The chat pipeline is fundamentally text + (optionally) vision. So an upload is
turned into one of:

* a **document** — plain text extracted from PDF / Word / Excel / PowerPoint /
  CSV / markdown / HTML / JSON / source files. Attached to the message as text.
* an **image** — returned as a base64 data URL for vision-capable models, plus
  best-effort OCR text as a fallback for models without vision.

Anything a text/vision agent genuinely can't use (archives, audio, video,
executables, unknown binaries) is rejected with a clear message. Files are
never persisted — we read the bytes, extract, and discard.
"""
import base64
import io
import json as _json
import logging

from fastapi import HTTPException, UploadFile, status
from starlette.concurrency import run_in_threadpool

logger = logging.getLogger("backend.attachments")

MAX_TEXT_CHARS = 200_000
MAX_DOC_BYTES = 20 * 1024 * 1024   # 20 MB for documents
MAX_IMAGE_BYTES = 10 * 1024 * 1024  # 10 MB for images
MAX_PDF_PAGES = 100
# If the extracted text layer averages fewer chars/page than this, the PDF is
# probably scanned/image-based (e.g. slide screenshots) — fall back to OCR.
MIN_CHARS_PER_PAGE = 20
PDF_OCR_ZOOM = 2.0  # ~144 DPI render for OCR — enough for legible text.

# Magic-byte signatures — the client filename/Content-Type are hints, never trusted.
_PDF_SIG = b"%PDF-"
_ZIP_SIG = b"PK\x03\x04"  # docx/xlsx/pptx are all zip containers

# Image signatures → mime type.
_IMAGE_SIGS: list[tuple[bytes, str]] = [
    (b"\x89PNG\r\n\x1a\n", "image/png"),
    (b"\xff\xd8\xff", "image/jpeg"),
    (b"GIF87a", "image/gif"),
    (b"GIF89a", "image/gif"),
    (b"BM", "image/bmp"),
]

# Plain-text-ish extensions we decode directly (no signature to sniff).
_TEXT_EXTS = (
    ".txt", ".md", ".markdown", ".text", ".csv", ".tsv", ".log",
    ".html", ".htm", ".xml", ".json", ".yaml", ".yml",
    ".py", ".js", ".ts", ".tsx", ".jsx", ".java", ".c", ".cpp", ".h", ".hpp",
    ".cs", ".go", ".rs", ".rb", ".php", ".sql", ".sh", ".css", ".ini", ".toml",
)


def _sniff_image_mime(data: bytes, name: str) -> str | None:
    for sig, mime in _IMAGE_SIGS:
        if data.startswith(sig):
            return mime
    if len(data) >= 12 and data[0:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "image/webp"
    # TIFF (little/big endian)
    if data[:4] in (b"II*\x00", b"MM\x00*"):
        return "image/tiff"
    return None


# ── Document extractors (all synchronous / CPU-bound; run off the event loop) ──

def _ocr_pdf_page(page) -> str:
    """Render a page to an image and OCR it. Returns '' if tesseract is unavailable."""
    try:
        import pytesseract
        from PIL import Image

        pix = page.get_pixmap(matrix=_fitz_matrix(PDF_OCR_ZOOM, PDF_OCR_ZOOM))
        img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        return pytesseract.image_to_string(img).strip()
    except Exception as exc:
        logger.info("PDF page OCR unavailable/failed: %s", exc)
        return ""


def _fitz_matrix(x: float, y: float):
    import fitz

    return fitz.Matrix(x, y)


def _extract_pdf(data: bytes) -> str:
    import fitz  # PyMuPDF — C-backed, fast even on large/image-heavy PDFs.

    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Could not read the PDF — it may be corrupt.")
    try:
        if doc.needs_pass:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, "Password-protected PDFs are not supported.")
        page_count = min(doc.page_count, MAX_PDF_PAGES)
        parts = [doc[i].get_text() for i in range(page_count)]

        # Scanned/screenshot-based PDFs (e.g. slide decks) have little or no
        # text layer — fall back to OCR per page so the agent still gets content.
        total_chars = sum(len(p.strip()) for p in parts)
        if page_count and total_chars < MIN_CHARS_PER_PAGE * page_count:
            ocr_parts = [_ocr_pdf_page(doc[i]) for i in range(page_count)]
            if sum(len(p) for p in ocr_parts) > total_chars:
                parts = ocr_parts
    finally:
        doc.close()
    return "\n\n".join(parts).strip()


def _extract_docx(data: bytes) -> str:
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
    except Exception:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Could not read the Word document.")
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def _extract_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Could not read the spreadsheet.")
    out: list[str] = []
    for ws in wb.worksheets:
        out.append(f"# Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                out.append(" | ".join(cells))
    wb.close()
    return "\n".join(out).strip()


def _extract_pptx(data: bytes) -> str:
    from pptx import Presentation

    try:
        prs = Presentation(io.BytesIO(data))
    except Exception:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Could not read the presentation.")
    out: list[str] = []
    for i, slide in enumerate(prs.slides, 1):
        out.append(f"# Slide {i}")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append(shape.text_frame.text.strip())
    return "\n".join(out).strip()


def _extract_text_plain(data: bytes) -> str:
    try:
        return data.decode("utf-8").strip()
    except UnicodeDecodeError:
        return data.decode("latin-1", errors="replace").strip()


def _ocr_image(data: bytes) -> str:
    """Best-effort OCR. Returns '' if the tesseract binary isn't installed —
    vision-capable models don't need it, and we degrade rather than 500."""
    try:
        import pytesseract
        from PIL import Image

        return pytesseract.image_to_string(Image.open(io.BytesIO(data))).strip()
    except Exception as exc:  # tesseract missing, unreadable image, etc.
        logger.info("OCR unavailable/failed: %s", exc)
        return ""


def _truncate(text: str) -> tuple[str, bool]:
    if len(text) > MAX_TEXT_CHARS:
        return text[:MAX_TEXT_CHARS], True
    return text, False


async def extract_document_text(data: bytes, filename: str) -> str:
    """Dispatch raw bytes to the right extractor by signature/extension.

    Shared by process_attachment (chat attachments) and any caller that
    stores raw bytes (e.g. GridFS) and needs to re-extract text later —
    keeps the sniffing/dispatch logic in one place.
    """
    name = (filename or "").lower()
    if data.startswith(_PDF_SIG):
        text = await run_in_threadpool(_extract_pdf, data)
    elif data.startswith(_ZIP_SIG) and name.endswith(".docx"):
        text = await run_in_threadpool(_extract_docx, data)
    elif data.startswith(_ZIP_SIG) and name.endswith(".xlsx"):
        text = await run_in_threadpool(_extract_xlsx, data)
    elif data.startswith(_ZIP_SIG) and name.endswith(".pptx"):
        text = await run_in_threadpool(_extract_pptx, data)
    elif name.endswith(_TEXT_EXTS):
        text = _extract_text_plain(data)
    else:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            "Unsupported file type. Upload an image, PDF, Word/Excel/PowerPoint, "
            "or a text/markdown/CSV/code file.",
        )
    return text


async def process_attachment(file: UploadFile) -> dict:
    """Turn an upload into a document/image result, or raise HTTPException(400).

    Returns either:
      {"kind": "document", "filename", "text", "chars", "truncated"}
      {"kind": "image", "filename", "mime", "data_url", "ocr_text"}
    """
    data = await file.read()
    if not data:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Uploaded file is empty.")

    name = (file.filename or "").lower()

    # ── Images ──────────────────────────────────────────────────────────────
    image_mime = _sniff_image_mime(data, name)
    if image_mime:
        if len(data) > MAX_IMAGE_BYTES:
            raise HTTPException(
                status.HTTP_400_BAD_REQUEST,
                f"Image too large. Max size is {MAX_IMAGE_BYTES // (1024 * 1024)}MB.",
            )
        ocr_text = await run_in_threadpool(_ocr_image, data)
        b64 = base64.b64encode(data).decode("ascii")
        return {
            "kind": "image",
            "filename": file.filename or "image",
            "mime": image_mime,
            "data_url": f"data:{image_mime};base64,{b64}",
            "ocr_text": ocr_text[:MAX_TEXT_CHARS],
        }

    # ── Documents ───────────────────────────────────────────────────────────
    if len(data) > MAX_DOC_BYTES:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            f"File too large. Max size is {MAX_DOC_BYTES // (1024 * 1024)}MB.",
        )

    text = await extract_document_text(data, file.filename or "")

    if not text:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            "No readable text found in the document (it may be a scanned image — "
            "upload it as an image instead).",
        )

    text, truncated = _truncate(text)
    return {
        "kind": "document",
        "filename": file.filename or "document",
        "text": text,
        "chars": len(text),
        "truncated": truncated,
    }
